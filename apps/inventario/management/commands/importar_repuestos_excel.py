from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.inventario.models import (
    Categoria,
    InventarioStock,
    MarcaRepuesto,
    MovimientoInventario,
    Repuesto,
    UbicacionFisica,
    UnidadMedida,
)


GENERIC_BRAND = "Sin Marca"
INITIAL_LOCATION = "STOCK-INICIAL-EXCEL"

BRAND_ALIASES = {
    "FLEETGUARTD": "FLEETGUARD",
    "FMX VALVES SETS": "FMX",
    "HINO ( ORIGINAL )": "HINO",
    "HYUNDAI ( ORIGINAL )": "HYUNDAI",
    "HYUNDAI ACCENT": "HYUNDAI",
    "HYUNDAI ELANTRA": "HYUNDAI",
    "HYUNDAI SANTA FE": "HYUNDAI",
    "IHP FILTER": "IHP FILTERS",
    "IHPFILTER": "IHP FILTERS",
    "IHPFILTERS": "IHP FILTERS",
    "MITSUBISHI ( ORIGINAL )": "MITSUBISHI",
    "PLUSPARTS": "PLUS PARTS",
    "SUZUKI GRAND VITARA": "SUZUKI",
    "TOYOTA ( ORIGINAL )": "TOYOTA",
    "TOYOTA ( ORRIGINAL )": "TOYOTA",
    "VISTONY (LOTOX)": "VISTONY",
}

GENERIC_BRAND_VALUES = {
    "",
    "AIR FILTER",
    "ALTERNATIVO",
    "CHERVOLET SAIL",
    "FILTERS",
    "GENUINE",
    "GENUINE PART",
    "GENUINE PARTS",
    "GENUINEPARTS",
    "POWER STEERING PUMP",
    "SIN MARCA",
    "USE FO JAC",
    "USE FORD DONGFENG",
    "USE FORD HYUNDAI",
}


@dataclass
class ImportRow:
    excel_row: int
    nombre: str
    codigo: str
    codigo_barra: str | None
    marca_nombre: str
    categoria_nombre: str
    cantidad: Decimal


def normalize(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).strip()).upper()


def classify_category(nombre):
    if nombre.startswith(("FILTRO", "FITRO")):
        return "Filtros"
    if nombre.startswith("ACEITE") or nombre.startswith("POTE DE GRASA"):
        return "Aceites y Lubricantes"
    if nombre.startswith(("AGUA PARA BATERIA", "LIMPIA PARABRISAS", "LIQUIDO DE FRENO", "LIQUIDO REFRIGERANTE")):
        return "Fluidos Vehiculares"
    if nombre.startswith(("BOMBA DE AGUA", "JUEGO DE PISTONES", "JUEGO DE VALVULA", "POLEA")):
        return "Motor y Distribucion"
    if nombre.startswith("BOMBA DE DIRECCON"):
        return "Direccion Hidraulica"
    if nombre.startswith("BOBINA DE ENCENDIDO"):
        return "Sistema Electrico y Encendido"
    if nombre.startswith(("PROTECTOR", "SILICONA", "SPAY")):
        return "Proteccion y Limpieza"
    if nombre.startswith("VALVULA DE PROTECCION"):
        return "Sistema Neumatico y Frenos"
    raise CommandError(f"No se encontro categoria para el repuesto '{nombre}'.")


def parse_quantity(value):
    text = normalize(value)
    if not text:
        return Decimal("0")
    match = re.search(r"\d+(?:[.,]\d+)?", text)
    if not match:
        raise CommandError(f"Cantidad invalida: '{text}'.")
    try:
        return Decimal(match.group().replace(",", "."))
    except InvalidOperation as error:
        raise CommandError(f"Cantidad invalida: '{text}'.") from error


def unique_code(base_code, excel_row, repeated_codes):
    if not base_code:
        return f"SIN-COD-{excel_row:03d}"
    if base_code not in repeated_codes:
        return base_code
    suffix = f"-F{excel_row}"
    return f"{base_code[:50 - len(suffix)]}{suffix}"


class Command(BaseCommand):
    help = "Importa repuestos y stock inicial desde el inventario Excel."

    def add_arguments(self, parser):
        parser.add_argument("--archivo", required=True, help="Ruta absoluta del archivo .xlsx")
        parser.add_argument(
            "--confirmar",
            action="store_true",
            help="Confirma la escritura. Sin esta opcion solo realiza una simulacion.",
        )

    def handle(self, *args, **options):
        path = Path(options["archivo"])
        if not path.is_file():
            raise CommandError(f"No se encontro el archivo: {path}")

        rows = self._read_rows(path)
        self._validate_database_dependencies(rows)
        summary = self._summary(rows)
        self.stdout.write(
            "Prevalidacion: "
            f"{summary['total']} repuestos, "
            f"{summary['generated_codes']} codigos generados, "
            f"{summary['suffixed_codes']} codigos repetidos diferenciados, "
            f"{summary['generic_brands']} marcas asignadas a Sin Marca, "
            f"{summary['blank_barcodes']} codigos de barra vacios por ausencia o repeticion, "
            f"{summary['zero_quantity']} cantidades iniciales en cero."
        )

        if not options["confirmar"]:
            self.stdout.write(self.style.WARNING(
                "Simulacion correcta. Use --confirmar para ejecutar la importacion."
            ))
            return

        self._import_rows(rows, path.name)
        self.stdout.write(self.style.SUCCESS(
            f"Importacion completada: {len(rows)} repuestos creados en {INITIAL_LOCATION}."
        ))

    def _read_rows(self, path):
        workbook = load_workbook(path, data_only=True, read_only=True)
        worksheet = workbook.active
        source_rows = []

        for excel_row, values in enumerate(worksheet.iter_rows(min_row=3, values_only=True), start=3):
            nombre = normalize(values[0])
            if not nombre:
                continue
            source_rows.append({
                "excel_row": excel_row,
                "nombre": nombre,
                "codigo": normalize(values[1]),
                "codigo_barra": normalize(values[2]),
                "marca": normalize(values[3]),
                "cantidad": parse_quantity(values[4]),
            })

        if not source_rows:
            raise CommandError("El archivo no contiene repuestos para importar.")

        repeated_codes = {
            code for code, count in Counter(row["codigo"] for row in source_rows if row["codigo"]).items()
            if count > 1
        }
        seen_barcodes = set()
        rows = []

        for source in source_rows:
            source_brand = source["marca"]
            normalized_brand = BRAND_ALIASES.get(source_brand, source_brand)
            if normalized_brand in GENERIC_BRAND_VALUES:
                normalized_brand = GENERIC_BRAND

            barcode = source["codigo_barra"] or None
            if barcode and barcode in seen_barcodes:
                barcode = None
            elif barcode:
                seen_barcodes.add(barcode)

            rows.append(ImportRow(
                excel_row=source["excel_row"],
                nombre=source["nombre"],
                codigo=unique_code(source["codigo"], source["excel_row"], repeated_codes),
                codigo_barra=barcode,
                marca_nombre=normalized_brand,
                categoria_nombre=classify_category(source["nombre"]),
                cantidad=source["cantidad"],
            ))

        codes = [row.codigo for row in rows]
        if len(codes) != len(set(codes)):
            raise CommandError("La normalizacion genero codigos duplicados.")
        return rows

    def _validate_database_dependencies(self, rows):
        category_names = {row.categoria_nombre for row in rows}
        brand_names = {row.marca_nombre for row in rows}
        categories = Categoria.objects.filter(nombre__in=category_names)
        brands = MarcaRepuesto.objects.filter(estado=True)
        missing_categories = category_names - set(categories.values_list("nombre", flat=True))
        available_brands = {name.upper() for name in brands.values_list("nombre", flat=True)}
        missing_brands = {
            brand_name for brand_name in brand_names
            if brand_name.upper() not in available_brands
        }
        if missing_categories:
            raise CommandError(f"Categorias faltantes: {', '.join(sorted(missing_categories))}")
        if missing_brands:
            raise CommandError(f"Marcas faltantes: {', '.join(sorted(missing_brands))}")
        if not UnidadMedida.objects.filter(nombre__iexact="Unidad", estado=True).exists():
            raise CommandError("No se encontro una unidad de medida activa llamada Unidad.")
        if not UbicacionFisica.objects.filter(codigo=INITIAL_LOCATION).exists():
            raise CommandError(f"No se encontro la ubicacion {INITIAL_LOCATION}.")

        existing_codes = set(Repuesto.objects.filter(codigo__in=[row.codigo for row in rows]).values_list("codigo", flat=True))
        if existing_codes:
            preview = ", ".join(sorted(existing_codes)[:10])
            raise CommandError(f"Ya existen codigos que chocan con la importacion: {preview}")

        existing_barcodes = {
            barcode for barcode in Repuesto.objects.filter(
                codigo_barra__in=[row.codigo_barra for row in rows if row.codigo_barra]
            ).values_list("codigo_barra", flat=True)
        }
        if existing_barcodes:
            preview = ", ".join(sorted(existing_barcodes)[:10])
            raise CommandError(f"Ya existen codigos de barra que chocan con la importacion: {preview}")

    def _summary(self, rows):
        return {
            "total": len(rows),
            "generated_codes": sum(row.codigo.startswith("SIN-COD-") for row in rows),
            "suffixed_codes": sum("-F" in row.codigo for row in rows),
            "generic_brands": sum(row.marca_nombre == GENERIC_BRAND for row in rows),
            "blank_barcodes": sum(row.codigo_barra is None for row in rows),
            "zero_quantity": sum(row.cantidad == 0 for row in rows),
        }

    @transaction.atomic
    def _import_rows(self, rows, source_name):
        categories = {category.nombre: category for category in Categoria.objects.all()}
        brands = {brand.nombre.upper(): brand for brand in MarcaRepuesto.objects.filter(estado=True)}
        unit = UnidadMedida.objects.get(nombre__iexact="Unidad", estado=True)
        location = UbicacionFisica.objects.get(codigo=INITIAL_LOCATION)

        for row in rows:
            repuesto = Repuesto.objects.create(
                codigo=row.codigo,
                codigo_barra=row.codigo_barra,
                nombre=row.nombre,
                categoria=categories[row.categoria_nombre],
                marca=brands[row.marca_nombre.upper()],
                unidad_medida=unit,
                precio_compra=Decimal("0.00"),
                precio_por_mayor=Decimal("0.00"),
                precio_cash=Decimal("0.00"),
                precio_lista=Decimal("0.00"),
            )
            InventarioStock.objects.create(
                repuesto=repuesto,
                ubicacion=location,
                stock_disponible=row.cantidad,
            )
            if row.cantidad > 0:
                MovimientoInventario.objects.create(
                    repuesto=repuesto,
                    ubicacion=location,
                    tipo_movimiento=MovimientoInventario.TipoMovimiento.INVENTARIO_INICIAL,
                    cantidad=row.cantidad,
                    stock_resultante=row.cantidad,
                    motivo=f"Inventario inicial migrado desde {source_name}, fila {row.excel_row}",
                )
