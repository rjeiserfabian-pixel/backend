import io
from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.utils import timezone

from .models import (
    AsignacionHerramienta,
    Herramienta,
    IncidenciaHerramienta,
    PlanMantenimientoHerramienta,
    RegistroMantenimiento,
)
from .test_mantenimientos import FaseTresBase

RESUMEN = '/api/herramientas/resumen/'
REP = '/api/herramientas/reportes/'


class ResumenTests(FaseTresBase):
    def test_contadores_por_estado_y_valor(self):
        Herramienta.objects.filter(pk=self.herramienta.pk).update(costo_adquisicion=Decimal('100'))
        Herramienta.objects.create(
            nombre='Sierra', categoria=self.categoria, sucursal=self.sucursal, creado_por=self.admin,
            costo_adquisicion=Decimal('250.50'), estado_operativo='EN_REPARACION',
        )
        Herramienta.objects.create(
            nombre='Perdida', categoria=self.categoria, sucursal=self.sucursal, creado_por=self.admin,
            costo_adquisicion=Decimal('999'), estado_operativo='PERDIDA',
        )
        Herramienta.objects.create(
            nombre='Eliminada', categoria=self.categoria, sucursal=self.sucursal, creado_por=self.admin,
            estado=False,
        )
        resp = self.client.get(RESUMEN)
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data['total'], 3)
        self.assertEqual(resp.data['por_estado']['DISPONIBLE'], 1)
        self.assertEqual(resp.data['por_estado']['EN_REPARACION'], 1)
        self.assertEqual(resp.data['por_estado']['PERDIDA'], 1)
        self.assertEqual(resp.data['valor_total'], 350.5)  # las perdidas no suman

    def test_alertas_en_el_resumen(self):
        self.plan_vencido()
        self.crear_plan(nombre='Pronto', intervalo_dias=3)
        self.client.post(
            '/api/herramientas/incidencias/',
            {'herramienta': self.herramienta.id, 'tipo': 'DANO', 'descripcion': 'x'}, format='json',
        )
        r = self.client.post(
            '/api/herramientas/asignaciones/',
            {'herramienta': self.herramienta.id, 'tecnico': str(self.tecnico.pk), 'forzar': True,
             'motivo_excepcion': 'prueba'}, format='json',
        )
        AsignacionHerramienta.objects.filter(pk=r.data['id']).update(
            fecha_devolucion_esperada=timezone.localdate() - timedelta(days=2)
        )
        data = self.client.get(RESUMEN).data
        self.assertEqual(data['mantenimientos_vencidos'], 1)
        self.assertEqual(data['mantenimientos_por_vencer'], 1)
        self.assertEqual(data['prestamos_vencidos'], 1)
        self.assertEqual(data['incidencias_abiertas'], 1)
        self.assertEqual(self.client.get(RESUMEN, {'sucursal': self.otra_sucursal.id}).data['total'], 0)

    def test_sin_permisos_de_otros_modulos_los_contadores_vienen_nulos(self):
        usuario, c = self.cliente_con_permisos('solo_inventario', ['HERRAMIENTAS.INVENTARIO.VER'])
        data = c.get(RESUMEN).data
        self.assertEqual(data['total'], 1)
        for clave in ('valor_total', 'mantenimientos_vencidos', 'mantenimientos_por_vencer',
                      'prestamos_vencidos', 'incidencias_abiertas'):
            self.assertIsNone(data[clave], clave)

    def test_permiso_y_validacion(self):
        nadie = self._crear_usuario('nadie_resumen')
        from rest_framework.test import APIClient
        c = APIClient()
        c.force_authenticate(user=nadie)
        self.assertEqual(c.get(RESUMEN).status_code, 403)
        self.assertEqual(APIClient().get(RESUMEN).status_code, 401)
        self.assertEqual(self.client.get(RESUMEN, {'sucursal': 'abc'}).status_code, 400)


class ReportesTests(FaseTresBase):
    def json(self, tipo, **params):
        return self.client.get(f'{REP}{tipo}/', params)

    def test_inventario_con_costos_y_fila_total(self):
        Herramienta.objects.filter(pk=self.herramienta.pk).update(costo_adquisicion=Decimal('120.50'))
        resp = self.json('inventario')
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertIn('Costo de adquisicion', resp.data['headers'])
        self.assertEqual(resp.data['rows'][0][0], self.herramienta.codigo)
        self.assertEqual(resp.data['rows'][-1][0], 'TOTAL')
        self.assertEqual(resp.data['rows'][-1][-1], 120.5)

    def test_inventario_sin_permiso_de_costos_no_muestra_costos(self):
        Herramienta.objects.filter(pk=self.herramienta.pk).update(costo_adquisicion=Decimal('120.50'))
        usuario, c = self.cliente_con_permisos('lector_reportes', ['HERRAMIENTAS.REPORTES.VER'])
        resp = c.get(f'{REP}inventario/')
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn('Costo de adquisicion', resp.data['headers'])
        self.assertNotIn('120.5', str(resp.data['rows']))
        self.assertNotEqual(resp.data['rows'][-1][0], 'TOTAL')

    def test_costos_de_mantenimiento_exigen_permiso_de_costos(self):
        usuario, c = self.cliente_con_permisos('lector2', ['HERRAMIENTAS.REPORTES.VER'])
        self.assertEqual(c.get(f'{REP}costos-mantenimiento/').status_code, 403)
        usuario2, c2 = self.cliente_con_permisos(
            'lector3', ['HERRAMIENTAS.REPORTES.VER', 'HERRAMIENTAS.COSTOS.VER']
        )
        self.assertEqual(c2.get(f'{REP}costos-mantenimiento/').status_code, 200)

    def test_filtros_del_inventario(self):
        Herramienta.objects.create(
            nombre='Sierra', categoria=self.categoria, sucursal=self.otra_sucursal, creado_por=self.admin
        )
        self.assertEqual(len(self.json('inventario', sucursal=self.otra_sucursal.id).data['rows']) - 1, 1)
        self.assertEqual(len(self.json('inventario', estado_operativo='PERDIDA').data['rows']), 0)

    def test_por_responsable_muestra_atraso(self):
        r = self.entregar()
        AsignacionHerramienta.objects.filter(pk=r.data['id']).update(
            fecha_devolucion_esperada=timezone.localdate() - timedelta(days=5)
        )
        fila = self.json('por-responsable').data['rows'][0]
        self.assertEqual(fila[0], 'Tecnico1 Prueba')
        self.assertEqual(fila[-1], 5)

    def test_costos_de_mantenimiento_agrupa_y_suma(self):
        for costo in ('50', '70.25'):
            rid = self.client.post(
                '/api/herramientas/mantenimientos/',
                {'herramienta': self.herramienta.id, 'tipo': 'CORRECTIVO', 'descripcion': 'x'}, format='json',
            ).data['id']
            self.client.post(
                f'/api/herramientas/mantenimientos/{rid}/finalizar/',
                {'resultado': 'ok', 'estado_fisico_final': 'BUENO', 'costo': costo}, format='json',
            )
        data = self.json('costos-mantenimiento').data
        self.assertEqual(data['rows'][0][3:], [2, 0, 2, 120.25])
        self.assertEqual(data['rows'][-1][-1], 120.25)
        manana = (timezone.localdate() + timedelta(days=1)).isoformat()
        self.assertEqual(len(self.json('costos-mantenimiento', desde=manana).data['rows']), 0)

    def test_mas_fallas_cuenta_reparaciones_y_danos(self):
        self.client.post(
            '/api/herramientas/mantenimientos/',
            {'herramienta': self.herramienta.id, 'tipo': 'CORRECTIVO', 'descripcion': 'x'}, format='json',
        )
        for _ in range(2):
            self.client.post(
                '/api/herramientas/incidencias/',
                {'herramienta': self.herramienta.id, 'tipo': 'DANO', 'descripcion': 'golpe'}, format='json',
            )
        self.client.post(
            '/api/herramientas/incidencias/',
            {'herramienta': self.herramienta.id, 'tipo': 'ROBO', 'descripcion': 'no cuenta'}, format='json',
        )
        Herramienta.objects.create(
            nombre='Sin fallas', categoria=self.categoria, sucursal=self.sucursal, creado_por=self.admin
        )
        rows = self.json('mas-fallas').data['rows']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][4:], [1, 2, 3])

    def test_validaciones(self):
        self.assertEqual(self.json('no-existe').status_code, 404)
        self.assertEqual(self.json('inventario', desde='2026-02-01', hasta='2026-01-01').status_code, 400)
        self.assertEqual(self.json('inventario', desde='no-fecha').status_code, 400)
        self.assertEqual(self.json('inventario', formato='csv').status_code, 400)
        self.assertEqual(self.json('inventario', estado_operativo='XXX').status_code, 400)

    def test_vista_previa_se_recorta_pero_conserva_el_total(self):
        Herramienta.objects.bulk_create([
            Herramienta(codigo=f'HX-{i:04d}', nombre=f'T{i}', categoria=self.categoria,
                        sucursal=self.sucursal, costo_adquisicion=Decimal('1'))
            for i in range(250)
        ])
        data = self.json('inventario').data
        self.assertTrue(data['truncado'])
        self.assertEqual(data['total_filas'], 252)  # 251 herramientas + fila TOTAL
        self.assertEqual(len(data['rows']), 200)
        self.assertEqual(data['rows'][-1][0], 'TOTAL')

    def test_permisos_de_reportes(self):
        from rest_framework.test import APIClient
        nadie = self._crear_usuario('nadie_rep')
        c = APIClient()
        c.force_authenticate(user=nadie)
        self.assertEqual(c.get(f'{REP}inventario/').status_code, 403)
        self.assertEqual(APIClient().get(f'{REP}inventario/').status_code, 401)


class ExportacionTests(FaseTresBase):
    def test_excel_valido_y_neutraliza_formulas(self):
        Herramienta.objects.filter(pk=self.herramienta.pk).update(nombre='=HYPERLINK("http://x","click")')
        resp = self.client.get(f'{REP}inventario/', {'formato': 'excel'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertTrue(resp.content.startswith(b'PK'))
        hoja = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual(hoja.cell(row=1, column=1).value, 'Codigo')
        self.assertEqual(hoja.cell(row=2, column=2).value, '\'=HYPERLINK("http://x","click")')

    def test_pdf_valido_aun_con_marcado_en_los_datos(self):
        Herramienta.objects.filter(pk=self.herramienta.pk).update(nombre='<b>Amoladora</b> & <i>cia')
        resp = self.client.get(f'{REP}inventario/', {'formato': 'pdf'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_todos_los_reportes_se_exportan(self):
        for tipo in ('inventario', 'por-responsable', 'costos-mantenimiento', 'mas-fallas'):
            for formato in ('excel', 'pdf'):
                resp = self.client.get(f'{REP}{tipo}/', {'formato': formato})
                self.assertEqual(resp.status_code, 200, f'{tipo} {formato}')

    def test_exportar_exige_su_permiso(self):
        usuario, c = self.cliente_con_permisos('solo_ver_rep', ['HERRAMIENTAS.REPORTES.VER'])
        self.assertEqual(c.get(f'{REP}inventario/').status_code, 200)
        self.assertEqual(c.get(f'{REP}inventario/', {'formato': 'excel'}).status_code, 403)
        self.assertEqual(c.get(f'{REP}inventario/', {'formato': 'pdf'}).status_code, 403)
        usuario2, c2 = self.cliente_con_permisos(
            'exportador', ['HERRAMIENTAS.REPORTES.VER', 'HERRAMIENTAS.REPORTES.EXPORTAR']
        )
        self.assertEqual(c2.get(f'{REP}inventario/', {'formato': 'excel'}).status_code, 200)
